"""
load_metas_cintia.py — Processa a planilha BAse MEtas Cintia.xlsx e gera metas_cintia.json
com agregações hierárquicas (Total, Distritais, Coordenadores, Filiais, Grupos e Linhas).
"""
import os
import sys
import time
import json
import openpyxl

if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'): sys.stderr.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "BAse MEtas Cintia.xlsx")
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_JSON = os.path.join(DATA_DIR, "metas_cintia.json")

def clean_name(val):
    if not val:
        return ""
    return str(val).strip().replace('\xa0', ' ')

def load_metas():
    t0 = time.time()
    print("=" * 70)
    print("  CARREGAMENTO DE METAS DIRETORIA C (CÍNTIA SILVA)")
    print("=" * 70)

    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Arquivo não encontrado: {EXCEL_PATH}")

    os.makedirs(DATA_DIR, exist_ok=True)
    print(f"Lendo planilha: {EXCEL_PATH} ...", flush=True)

    # Carrega com data_only=True para ler valores calculados das fórmulas
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    ws = wb['Meta Diarizada']

    # Linha 1: Pesos diários
    row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    weights = []
    for c in range(7, 37):
        w = row1[c] if c < len(row1) else 0.0
        try:
            weights.append(float(w or 0.0))
        except:
            weights.append(0.0)

    # Se o peso do dia 30 estiver zerado ou ausente na planilha, ajusta como complemento de 100%
    if len(weights) >= 30 and (weights[29] == 0.0 or weights[29] is None):
        sum_29 = sum(weights[:29])
        weights[29] = round(max(0.0, 1.0 - sum_29), 6)
        print(f"Ajustado peso do Dia 30: {weights[29]:.4f} (Complemento 100% da curva)")

    # Linha 2: Cabeçalho
    row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    days_header = []
    for c in range(7, 37):
        val = row2[c] if c < len(row2) else f"Dia {c-6}"
        if hasattr(val, 'strftime'):
            days_header.append(val.strftime('%d/%m/%Y'))
        else:
            days_header.append(f"{c-6:02d}/09/2026")

    print(f"Pesos diários identificados para {len(weights)} dias. Soma dos pesos: {sum(weights):.4f}")

    # Estruturas de agregação
    total_meta_mes = 0.0
    total_dias = [0.0] * 30

    distritais_data = {}
    coordenadores_data = {}
    filiais_data = {}
    grupos_data = {}

    row_count = 0
    print("Processando linhas da planilha (iterador rápido)...", flush=True)

    # Iterando linhas a partir da linha 3
    for r in ws.iter_rows(min_row=3, values_only=True):
        row_count += 1
        if row_count % 25000 == 0:
            print(f"   ... {row_count:,} linhas processadas", flush=True)

        distrital = clean_name(r[1])
        coordenador = clean_name(r[2])
        filial = clean_name(r[3])
        grupo = clean_name(r[4])
        linha = clean_name(r[5])
        meta_val = r[6]

        try:
            meta_mes = float(meta_val or 0.0)
        except:
            meta_mes = 0.0

        if meta_mes == 0.0 and not distrital:
            continue

        # Metas diarizadas desta linha
        metas_dias_row = []
        for d_idx in range(30):
            col_val = r[7 + d_idx] if (7 + d_idx) < len(r) else None
            # Se for o dia 30 e o valor da célula for 0 ou nulo, aplica o peso do dia 30
            if d_idx == 29 and (col_val is None or col_val == 0 or col_val == 0.0):
                val = meta_mes * weights[29]
            else:
                try:
                    val = float(col_val) if col_val is not None else (meta_mes * weights[d_idx])
                except:
                    val = meta_mes * weights[d_idx]
            metas_dias_row.append(val)

        # 1. Total Geral
        total_meta_mes += meta_mes
        for d_idx in range(30):
            total_dias[d_idx] += metas_dias_row[d_idx]

        # 2. Distrital
        if distrital:
            if distrital not in distritais_data:
                distritais_data[distrital] = {
                    "meta_mes": 0.0,
                    "metas_dias": [0.0] * 30,
                    "coordenadores": set(),
                    "filiais": set()
                }
            distritais_data[distrital]["meta_mes"] += meta_mes
            for d_idx in range(30):
                distritais_data[distrital]["metas_dias"][d_idx] += metas_dias_row[d_idx]
            if coordenador:
                distritais_data[distrital]["coordenadores"].add(coordenador)
            if filial:
                distritais_data[distrital]["filiais"].add(filial)

        # 3. Coordenador
        if coordenador:
            if coordenador not in coordenadores_data:
                coordenadores_data[coordenador] = {
                    "distrital": distrital,
                    "meta_mes": 0.0,
                    "metas_dias": [0.0] * 30,
                    "filiais": set()
                }
            coordenadores_data[coordenador]["meta_mes"] += meta_mes
            for d_idx in range(30):
                coordenadores_data[coordenador]["metas_dias"][d_idx] += metas_dias_row[d_idx]
            if filial:
                coordenadores_data[coordenador]["filiais"].add(filial)

        # 4. Filial
        if filial:
            if filial not in filiais_data:
                filiais_data[filial] = {
                    "distrital": distrital,
                    "coordenador": coordenador,
                    "meta_mes": 0.0,
                    "metas_dias": [0.0] * 30
                }
            filiais_data[filial]["meta_mes"] += meta_mes
            for d_idx in range(30):
                filiais_data[filial]["metas_dias"][d_idx] += metas_dias_row[d_idx]

        # 5. Grupo & Linha
        if grupo:
            if grupo not in grupos_data:
                grupos_data[grupo] = {
                    "meta_mes": 0.0,
                    "metas_dias": [0.0] * 30,
                    "linhas": {}
                }
            grupos_data[grupo]["meta_mes"] += meta_mes
            for d_idx in range(30):
                grupos_data[grupo]["metas_dias"][d_idx] += metas_dias_row[d_idx]

            if linha:
                if linha not in grupos_data[grupo]["linhas"]:
                    grupos_data[grupo]["linhas"][linha] = {
                        "meta_mes": 0.0,
                        "metas_dias": [0.0] * 30
                    }
                grupos_data[grupo]["linhas"][linha]["meta_mes"] += meta_mes
                for d_idx in range(30):
                    grupos_data[grupo]["linhas"][linha]["metas_dias"][d_idx] += metas_dias_row[d_idx]

    # Converte sets em listas para JSON serializável e arredonda
    for d, info in distritais_data.items():
        info["coordenadores"] = sorted(list(info["coordenadores"]))
        info["filiais"] = sorted(list(info["filiais"]))
        info["meta_mes"] = round(info["meta_mes"], 2)
        info["metas_dias"] = [round(v, 2) for v in info["metas_dias"]]

    for c, info in coordenadores_data.items():
        info["filiais"] = sorted(list(info["filiais"]))
        info["meta_mes"] = round(info["meta_mes"], 2)
        info["metas_dias"] = [round(v, 2) for v in info["metas_dias"]]

    for f, info in filiais_data.items():
        info["meta_mes"] = round(info["meta_mes"], 2)
        info["metas_dias"] = [round(v, 2) for v in info["metas_dias"]]

    for g, info in grupos_data.items():
        info["meta_mes"] = round(info["meta_mes"], 2)
        info["metas_dias"] = [round(v, 2) for v in info["metas_dias"]]
        for l, l_info in info["linhas"].items():
            l_info["meta_mes"] = round(l_info["meta_mes"], 2)
            l_info["metas_dias"] = [round(v, 2) for v in l_info["metas_dias"]]

    payload = {
        "metadata": {
            "gerado_em": time.strftime("%Y-%m-%d %H:%M:%S"),
            "mes": "Setembro/2026",
            "dias_totais": 30,
            "dias_header": days_header,
            "weights": [round(w, 4) for w in weights],
            "total_linhas_planilha": row_count,
            "total_filiais": len(filiais_data),
            "total_coordenadores": len(coordenadores_data),
            "total_distritais": len(distritais_data),
            "total_grupos": len(grupos_data)
        },
        "total": {
            "meta_mes": round(total_meta_mes, 2),
            "metas_dias": [round(v, 2) for v in total_dias]
        },
        "distritais": distritais_data,
        "coordenadores": coordenadores_data,
        "filiais": filiais_data,
        "grupos": grupos_data
    }

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    elapsed = time.time() - t0
    print(f"✅ Processamento concluído em {elapsed:.1f}s")
    print(f"   📊 Meta Total do Mês: R$ {total_meta_mes:,.2f}")
    print(f"   🏢 Distritais: {len(distritais_data)} | Coordenadores: {len(coordenadores_data)} | Filiais: {len(filiais_data)}")
    print(f"   📁 Salvo em: {OUTPUT_JSON} ({os.path.getsize(OUTPUT_JSON)/1024:.1f} KB)")
    return payload

if __name__ == "__main__":
    load_metas()
